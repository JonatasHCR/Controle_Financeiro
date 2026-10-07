"""Modelo de planilha da configuração: gerar (em branco ou preenchido) e importar.

Uma aba por grupo de complemento, coluna CR em todas: uma planilha pode trazer
vários contratos. Para cada CR presente numa aba, as linhas da aba passam a ser
o estado daquele grupo (o que não veio é removido) — menos BMs e de-para, onde
o que não veio fica como está.
"""

from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook
from sqlalchemy import select

from app.analise.dados import carregar_contratos
from app.analise.normalizar import cr_norm, nome_norm
from app.configuracao import gravar
from app.configuracao.servico import (
    ErroDeValidacao,
    data_br,
    decimal_br,
    mes_br,
    naturezas_do_cr,
    nfs_do_cr,
    opcao,
)
from app.models import (
    CdNatureza,
    CfgContrato,
    CfgItem,
    CfgPendencia,
    CfgPleito,
    Parametros,
    RcContrato,
    RcNf,
)
from app.models.complementos import PRIORIDADES, STATUS_PENDENCIA, STATUS_PLEITO, TIPOS_PLEITO

ABAS = {
    "Contrato": ["CR", "Nome (informativo)", "Fim da execução", "Tributos %", "Taxa adm. %"],
    "Itens": [
        "CR",
        "Código",
        "Descrição",
        "Unidade",
        "Quantidade",
        "Custo-alvo",
        "Custo-alvo projetado",
        "Quantidade medida",
    ],
    "De-para": ["CR", "Natureza (Controle de Despesa)", "Código do item"],
    "BMs": [
        "CR",
        "NF",
        "Emissão (informativo)",
        "Valor (informativo)",
        "BM",
        "Data-base (mm/aaaa)",
    ],
    "Pendências": ["CR", "Assunto", "Prazo", "Prioridade", "Status"],
    "Pleitos": ["CR", "Data-base (mm/aaaa)", "Descrição", "Tipo", "Status", "Valor", "Observação"],
}

INSTRUCOES = [
    "Modelo de configuração do Controle Financeiro (UFC Engenharia).",
    "",
    "Aqui só entram complementos: contratos, NFs e despesas vêm do Gerenciamento de Receita"
    " e do Controle de Despesa.",
    "Cada aba tem a coluna CR; uma planilha pode trazer vários contratos.",
    "Para cada CR que aparece numa aba, as linhas da aba passam a ser o estado daquele grupo:",
    "  • Itens, Pendências e Pleitos: o que não estiver na planilha é removido.",
    "  • De-para e BMs: o que não estiver na planilha fica como está; BM vazio tira o BM da NF.",
    "  • Contrato: tributos e taxa vazios voltam ao padrão global (20% e 15%).",
    "Datas: dd/mm/aaaa. Data-base: mm/aaaa. Valores: 1234,56.",
    "Colunas marcadas (informativo) são ignoradas na importação.",
    "A importação mostra uma prévia por aba antes de gravar.",
]


# --- gerar ------------------------------------------------------------------------


def gerar(session, crs: list[str]) -> io.BytesIO:
    saida = io.BytesIO()
    livro = xlsxwriter.Workbook(saida, {"in_memory": True})
    cab = livro.add_format(
        {"bold": True, "bg_color": "#A51C24", "font_color": "#FFFFFF", "border": 1}
    )
    data = livro.add_format({"num_format": "dd/mm/yyyy"})
    mes = livro.add_format({"num_format": "mm/yyyy"})
    dinheiro = livro.add_format({"num_format": "#,##0.00"})
    pct = livro.add_format({"num_format": "0.00"})
    info = livro.add_format({"font_color": "#857A7B", "italic": True})

    instrucoes = livro.add_worksheet("Instruções")
    instrucoes.set_column(0, 0, 120)
    for i, linha in enumerate(INSTRUCOES):
        instrucoes.write(i, 0, linha, livro.add_format({"bold": True}) if i == 0 else None)

    listas = livro.add_worksheet("Listas")
    listas.hide()
    for col, (nome, valores) in enumerate(
        (
            ("Prioridade", PRIORIDADES),
            ("Status pendência", STATUS_PENDENCIA),
            ("Tipo pleito", TIPOS_PLEITO),
            ("Status pleito", STATUS_PLEITO),
        )
    ):
        listas.write(0, col, nome)
        listas.write_column(1, col, valores)

    folhas = {}
    for nome, colunas in ABAS.items():
        folha = livro.add_worksheet(nome)
        folha.write_row(0, 0, colunas, cab)
        folha.freeze_panes(1, 0)
        for c, titulo in enumerate(colunas):
            folha.set_column(c, c, max(12, min(48, len(titulo) + 6)))
        folhas[nome] = folha

    def lista(folha, coluna, intervalo):
        folha.data_validation(1, coluna, 2000, coluna, {"validate": "list", "source": intervalo})

    lista(folhas["Pendências"], 3, "=Listas!$A$2:$A$4")
    lista(folhas["Pendências"], 4, "=Listas!$B$2:$B$3")
    lista(folhas["Pleitos"], 3, "=Listas!$C$2:$C$4")
    lista(folhas["Pleitos"], 4, "=Listas!$D$2:$D$5")
    folhas["Contrato"].data_validation(
        1,
        2,
        2000,
        2,
        {
            "validate": "date",
            "criteria": ">",
            "value": 1,
            "error_message": "Use uma data (dd/mm/aaaa).",
        },
    )

    contratos = {c.cr: c for c in carregar_contratos(session)}
    linhas = defaultdict(int)

    def proxima(nome):
        linhas[nome] += 1
        return linhas[nome]

    for cr in crs:
        c = contratos.get(cr)
        if c is None:
            continue
        cfg = session.get(CfgContrato, cr)
        par = session.get(Parametros, cr)
        r = proxima("Contrato")
        folhas["Contrato"].write(r, 0, cr)
        folhas["Contrato"].write(r, 1, c.descricao, info)
        if cfg and cfg.fim_execucao:
            folhas["Contrato"].write_datetime(r, 2, _dt(cfg.fim_execucao), data)
        if par:
            folhas["Contrato"].write_number(r, 3, float(par.tributos) * 100, pct)
            folhas["Contrato"].write_number(r, 4, float(par.taxa_adm) * 100, pct)

        for it in session.scalars(
            select(CfgItem).where(CfgItem.cr_norm == cr).order_by(CfgItem.ordem)
        ):
            r = proxima("Itens")
            folhas["Itens"].write_row(r, 0, [cr, it.codigo, it.descricao, it.unidade])
            for col, valor in (
                (4, it.quantidade),
                (5, it.custo_alvo),
                (6, it.custo_alvo_projetado),
                (7, it.quantidade_medida),
            ):
                if valor is not None:
                    folhas["Itens"].write_number(
                        r, col, float(valor), dinheiro if col in (5, 6) else None
                    )

        for n in naturezas_do_cr(session, cr):
            r = proxima("De-para")
            folhas["De-para"].write_row(r, 0, [cr, n["nome"], n["item"] or n["sugestao"]])

        for nf in nfs_do_cr(session, cr):
            r = proxima("BMs")
            folhas["BMs"].write_row(r, 0, [cr, nf["numero"]])
            folhas["BMs"].write_datetime(r, 2, _dt(nf["emitida_em"]), data)
            folhas["BMs"].write_number(r, 3, nf["valor"], dinheiro)
            if nf["bm"] is not None:
                folhas["BMs"].write_number(r, 4, nf["bm"])
            if nf["data_base"]:
                folhas["BMs"].write_datetime(r, 5, _dt(nf["data_base"]), mes)

        for p in session.scalars(select(CfgPendencia).where(CfgPendencia.cr_norm == cr)):
            r = proxima("Pendências")
            folhas["Pendências"].write(r, 0, cr)
            folhas["Pendências"].write(r, 1, p.assunto)
            if p.prazo:
                folhas["Pendências"].write_datetime(r, 2, _dt(p.prazo), data)
            folhas["Pendências"].write_row(r, 3, [p.prioridade, p.status])

        for p in session.scalars(select(CfgPleito).where(CfgPleito.cr_norm == cr)):
            r = proxima("Pleitos")
            folhas["Pleitos"].write(r, 0, cr)
            if p.data_base:
                folhas["Pleitos"].write_datetime(r, 1, _dt(p.data_base), mes)
            folhas["Pleitos"].write_row(r, 2, [p.descricao, p.tipo, p.status])
            folhas["Pleitos"].write_number(r, 5, float(p.valor or 0), dinheiro)
            folhas["Pleitos"].write(r, 6, p.observacao or "")

    livro.close()
    saida.seek(0)
    return saida


def _dt(d):
    from datetime import datetime

    return datetime(d.year, d.month, d.day)


# --- importar ---------------------------------------------------------------------


@dataclass
class Resultado:
    por_aba: dict[str, gravar.Contagem] = field(
        default_factory=lambda: defaultdict(gravar.Contagem)
    )
    erros: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    crs: list[str] = field(default_factory=list)

    def somar(self, aba: str, contagem: gravar.Contagem) -> None:
        self.por_aba[aba] = self.por_aba[aba].somar(contagem)
        self.avisos.extend(contagem.avisos)

    def totais(self) -> gravar.Contagem:
        total = gravar.Contagem()
        for c in self.por_aba.values():
            total = total.somar(gravar.Contagem(c.novas, c.alteradas, c.removidas))
        return total

    def contagens_dict(self) -> dict:
        return {aba: c.como_dict() for aba, c in self.por_aba.items()}


def _linhas(folha, colunas: list[str]) -> list[tuple[int, dict]]:
    """Linhas não vazias com o cabeçalho esperado, como (número da linha, dict)."""
    cabecalho = [
        nome_norm(c.value) if c.value is not None else "" for c in next(folha.iter_rows(max_row=1))
    ]
    indices = {}
    for titulo in colunas:
        alvo = nome_norm(titulo)
        if alvo in cabecalho:
            indices[titulo] = cabecalho.index(alvo)
    saida = []
    for numero, linha in enumerate(folha.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v not in (None, "") for v in linha):
            continue
        saida.append(
            (numero, {t: (linha[i] if i < len(linha) else None) for t, i in indices.items()})
        )
    return saida


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def aplicar(session, caminho: Path, usuario_id: int | None, simular: bool) -> Resultado:
    r = Resultado()
    try:
        livro = load_workbook(caminho, data_only=True, read_only=True)
    except Exception:  # arquivo corrompido ou não é xlsx
        r.erros.append("Não foi possível abrir a planilha. Use o modelo .xlsx baixado do sistema.")
        return r
    contratos = {c.cr: c for c in carregar_contratos(session)}
    naturezas = set(session.scalars(select(CdNatureza.nome_norm)))
    tocados: set[str] = set()

    def grupos(aba: str) -> dict[str, list[tuple[int, dict]]]:
        if aba not in livro.sheetnames:
            return {}
        saida = defaultdict(list)
        for numero, linha in _linhas(livro[aba], ABAS[aba]):
            cr = cr_norm(_texto(linha.get("CR")))
            if not cr:
                r.erros.append(f"{aba}, linha {numero}: CR vazio")
                continue
            c = contratos.get(cr)
            if c is None:
                r.erros.append(f"{aba}, linha {numero}: CR {cr} não tem contrato na Receita")
                continue
            if not c.receita and aba != "De-para":
                r.erros.append(
                    f"{aba}, linha {numero}: CR {cr} só existe no Controle de Despesa"
                    " (só o de-para vale)"
                )
                continue
            saida[cr].append((numero, linha))
        return saida

    def com_erros(aba, cr, funcao):
        """Converte as linhas; uma linha errada barra o grupo inteiro daquele CR."""
        convertidas, ok = [], True
        for numero, linha in grupos_por_aba[aba][cr]:
            try:
                convertidas.append(funcao(linha))
            except (ErroDeValidacao, ValueError) as erro:
                r.erros.append(f"{aba}, linha {numero}: {erro}")
                ok = False
        return convertidas if ok else None

    grupos_por_aba = {aba: grupos(aba) for aba in ABAS}

    for cr, linhas in grupos_por_aba["Contrato"].items():
        numero, linha = linhas[-1]
        try:
            fim = data_br(linha.get("Fim da execução"), "Fim da execução")
            trib = decimal_br(linha.get("Tributos %"), "Tributos %")
            taxa = decimal_br(linha.get("Taxa adm. %"), "Taxa adm. %")
            r.somar("Contrato", gravar.salvar_prazos(session, cr, fim, usuario_id, simular))
            r.somar(
                "Contrato",
                gravar.salvar_parametros(
                    session,
                    cr,
                    trib / 100 if trib is not None else None,
                    taxa / 100 if taxa is not None else None,
                    usuario_id,
                    simular,
                ),
            )
            tocados.add(cr)
        except (ErroDeValidacao, ValueError) as erro:
            r.erros.append(f"Contrato, linha {numero}: {erro}")

    itens_da_planilha = defaultdict(set)
    for cr in grupos_por_aba["Itens"]:

        def item(linha):
            codigo, descricao = _texto(linha.get("Código")), _texto(linha.get("Descrição"))
            if not codigo or not descricao:
                raise ErroDeValidacao("código e descrição são obrigatórios")
            return {
                "codigo": codigo[:20],
                "descricao": descricao[:300],
                "unidade": _texto(linha.get("Unidade"))[:20],
                "quantidade": decimal_br(linha.get("Quantidade"), "quantidade"),
                "custo_alvo": decimal_br(linha.get("Custo-alvo"), "custo-alvo", obrigatorio=True),
                "custo_alvo_projetado": decimal_br(
                    linha.get("Custo-alvo projetado"), "custo-alvo projetado"
                ),
                "quantidade_medida": decimal_br(
                    linha.get("Quantidade medida"), "quantidade medida"
                ),
            }

        convertidas = com_erros("Itens", cr, item)
        if convertidas is None:
            continue
        try:
            r.somar("Itens", gravar.substituir_itens(session, cr, convertidas, usuario_id, simular))
            itens_da_planilha[cr] = {i["codigo"] for i in convertidas}
            tocados.add(cr)
        except ValueError as erro:
            r.erros.append(f"Itens, CR {cr}: {erro}")

    for cr in grupos_por_aba["De-para"]:
        validos = itens_da_planilha.get(cr) or set(
            session.scalars(select(CfgItem.codigo).where(CfgItem.cr_norm == cr))
        )
        mapa, ok = {}, True
        for numero, linha in grupos_por_aba["De-para"][cr]:
            natureza, item = (
                nome_norm(_texto(linha.get("Natureza (Controle de Despesa)"))),
                _texto(linha.get("Código do item")),
            )
            if not natureza:
                r.erros.append(f"De-para, linha {numero}: natureza vazia")
                ok = False
                continue
            if item and item not in validos:
                r.erros.append(
                    f"De-para, linha {numero}: item {item} não existe nos itens do CR {cr}"
                )
                ok = False
                continue
            nome_original = _texto(linha.get("Natureza (Controle de Despesa)"))
            if natureza not in naturezas:
                r.avisos.append(
                    f"De-para, linha {numero}: a natureza “{nome_original}”"
                    " ainda não existe no Controle de Despesa"
                )
            mapa[natureza] = item
        if ok:
            r.somar("De-para", gravar.salvar_depara(session, cr, mapa, usuario_id, simular))
            tocados.add(cr)

    for cr in grupos_por_aba["BMs"]:
        origem = session.scalar(select(RcContrato.origem_id).where(RcContrato.cr_norm == cr))
        nfs = {
            n.numero: n.origem_id
            for n in session.scalars(select(RcNf).where(RcNf.contrato_origem_id == origem))
        }
        linhas, ok = [], True
        for numero, linha in grupos_por_aba["BMs"][cr]:
            nf = _texto(linha.get("NF"))
            if nf not in nfs:
                r.avisos.append(
                    f"BMs, linha {numero}: a NF {nf} ainda não chegou da Receita;"
                    " a linha foi ignorada"
                )
                continue
            try:
                bm = decimal_br(linha.get("BM"), "BM")
                linhas.append(
                    {
                        "nf_origem_id": nfs[nf],
                        "numero": int(bm) if bm is not None else None,
                        "data_base": mes_br(linha.get("Data-base (mm/aaaa)"), "data-base"),
                    }
                )
            except ErroDeValidacao as erro:
                r.erros.append(f"BMs, linha {numero}: {erro}")
                ok = False
        if ok:
            r.somar("BMs", gravar.salvar_bms(session, cr, linhas, usuario_id, simular))
            tocados.add(cr)

    for cr in grupos_por_aba["Pendências"]:
        convertidas = com_erros(
            "Pendências",
            cr,
            lambda linha: {
                "assunto": _texto(linha.get("Assunto")) or _falta("assunto"),
                "prazo": data_br(linha.get("Prazo"), "prazo"),
                "prioridade": opcao(linha.get("Prioridade"), "prioridade", PRIORIDADES, "MÉDIA"),
                "status": opcao(linha.get("Status"), "status", STATUS_PENDENCIA, "ABERTA"),
            },
        )
        if convertidas is not None:
            r.somar(
                "Pendências",
                gravar.substituir_pendencias(session, cr, convertidas, usuario_id, simular),
            )
            tocados.add(cr)

    for cr in grupos_por_aba["Pleitos"]:
        convertidas = com_erros(
            "Pleitos",
            cr,
            lambda linha: {
                "data_base": mes_br(linha.get("Data-base (mm/aaaa)"), "data-base"),
                "descricao": _texto(linha.get("Descrição")) or _falta("descrição"),
                "tipo": opcao(linha.get("Tipo"), "tipo", TIPOS_PLEITO, "ADITIVO"),
                "status": opcao(linha.get("Status"), "status", STATUS_PLEITO, "SOLICITAR"),
                "valor": decimal_br(linha.get("Valor"), "valor"),
                "observacao": _texto(linha.get("Observação")) or None,
            },
        )
        if convertidas is not None:
            r.somar(
                "Pleitos", gravar.substituir_pleitos(session, cr, convertidas, usuario_id, simular)
            )
            tocados.add(cr)

    livro.close()
    r.crs = sorted(tocados, key=lambda c: (len(c), c))
    return r


def _falta(campo: str):
    raise ErroDeValidacao(f"{campo} é obrigatório")
