"""Configuração: só complementos, telas e planilha (ida e volta)."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import (
    Auditoria,
    CfgBm,
    CfgContrato,
    CfgItem,
    CfgPendencia,
    CfgPleito,
    DeparaItem,
    Parametros,
    RcContrato,
)

pytestmark = pytest.mark.integration


def test_leitor_nao_entra(entrar, leitor, carregado):
    assert entrar(leitor).get("/configuracao/").status_code == 403


def test_lista_mostra_os_contratos_da_receita(entrar, operador, carregado):
    corpo = entrar(operador).get("/configuracao/").get_data(as_text=True)
    for cr in ("4561", "4602", "4660"):
        assert cr in corpo
    assert "sem despesa" in corpo
    assert "4655" not in corpo


def test_nenhuma_rota_cria_contrato(app):
    regras = [r.rule for r in app.url_map.iter_rules() if r.rule.startswith("/configuracao")]
    assert not any("contrato" in r and "novo" in r for r in regras)


def test_prazo_de_execucao(entrar, operador, carregado, db):
    resposta = entrar(operador).post(
        "/configuracao/4561/prazos", data={"fim_execucao": "2027-03-31"}
    )
    assert resposta.status_code == 302
    assert db.session.get(CfgContrato, "4561").fim_execucao == date(2027, 3, 31)
    assert db.session.scalars(
        select(Auditoria).where(Auditoria.acao == "configuracao.prazos")
    ).one()


def test_itens_substituem_por_inteiro(entrar, operador, carregado, db):
    entrar(operador).post(
        "/configuracao/4561/itens",
        data={
            "it-0-codigo": "1.1",
            "it-0-descricao": "Engenheiro fiscal (PJ)",
            "it-0-custo_alvo": "450.000,00",
            "it-1-codigo": "1.2",
            "it-1-descricao": "Veículo locado",
            "it-1-custo_alvo": "100000",
            "it-1-remover": "on",
            "it-2-codigo": "1.3",
            "it-2-descricao": "Topografia",
            "it-2-custo_alvo": "25000,5",
            "it-3-codigo": "",
            "it-3-descricao": "",
        },
    )
    itens = {
        i.codigo: i.custo_alvo
        for i in db.session.scalars(select(CfgItem).where(CfgItem.cr_norm == "4561"))
    }
    assert itens == {"1.1": Decimal("450000.00"), "1.3": Decimal("25000.50")}


def test_item_sem_descricao_nao_salva_nada(entrar, operador, carregado, db):
    entrar(operador).post(
        "/configuracao/4561/itens", data={"it-0-codigo": "9.9", "it-0-descricao": ""}
    )
    assert db.session.scalars(select(CfgItem).where(CfgItem.codigo == "9.9")).first() is None
    assert len(db.session.scalars(select(CfgItem).where(CfgItem.cr_norm == "4561")).all()) == 2


def test_depara(entrar, operador, carregado, db):
    entrar(operador).post(
        "/configuracao/4561/depara",
        data={
            "dp-0-natureza": "LOCACAO DE VEICULOS",
            "dp-0-item": "",
        },
    )
    linha = db.session.scalars(
        select(DeparaItem).where(DeparaItem.natureza_nome_norm == "LOCACAO DE VEICULOS")
    ).one()
    assert linha.item_codigo == ""


def test_cr_so_do_controle_nao_tem_configuracao(entrar, operador, carregado):
    resposta = entrar(operador).post(
        "/configuracao/4655/depara", data={"dp-0-natureza": "LOCACAO DE VEICULOS", "dp-0-item": ""}
    )
    assert resposta.status_code == 404


def test_itens_mostram_as_naturezas_ligadas(entrar, operador, carregado):
    corpo = entrar(operador).get("/configuracao/?cr=4561&aba=itens").get_data(as_text=True)
    assert "Naturezas ligadas" in corpo


def test_depara_tem_a_propria_natureza_como_padrao(entrar, operador, carregado):
    corpo = entrar(operador).get("/configuracao/?cr=4602&aba=depara").get_data(as_text=True)
    assert "a própria natureza (padrão)" in corpo


def test_bms(entrar, operador, carregado, db):
    entrar(operador).post(
        "/configuracao/4561/bms",
        data={
            "bm-0-nf": "1",
            "bm-0-numero": "7",
            "bm-0-data_base": "2025-12",
            "bm-1-nf": "2",
            "bm-1-numero": "",
        },
    )
    bm = db.session.scalars(select(CfgBm)).one()
    assert (bm.nf_origem_id, bm.numero, bm.data_base) == (1, 7, date(2025, 12, 1))


def test_bm_muda_a_data_base_da_nf_no_painel(entrar, operador, leitor, carregado, db):
    entrar(operador).post(
        "/configuracao/4561/bms",
        data={"bm-0-nf": "1", "bm-0-numero": "7", "bm-0-data_base": "2025-12"},
    )
    dados = entrar(leitor).get("/api/painel?cr=4561&modo=mes&mes=2025-12").get_json()
    assert dados["cascata"]["fat"] == pytest.approx(100_000)


def test_pendencias_e_pleitos(entrar, operador, carregado, db):
    cliente = entrar(operador)
    cliente.post(
        "/configuracao/4561/pendencias",
        data={
            "pd-0-assunto": "Nova",
            "pd-0-prazo": "2026-05-01",
            "pd-0-prioridade": "BAIXA",
            "pd-0-status": "ABERTA",
        },
    )
    assert [
        p.assunto
        for p in db.session.scalars(select(CfgPendencia).where(CfgPendencia.cr_norm == "4561"))
    ] == ["Nova"]
    cliente.post(
        "/configuracao/4561/pleitos",
        data={
            "pl-0-descricao": "Reequilíbrio combustível",
            "pl-0-tipo": "REEQUILÍBRIO",
            "pl-0-status": "SOLICITAR",
            "pl-0-valor": "1.234,56",
            "pl-0-data_base": "2026-03",
        },
    )
    pleito = db.session.scalars(select(CfgPleito).where(CfgPleito.cr_norm == "4561")).one()
    assert (pleito.valor, pleito.data_base) == (Decimal("1234.56"), date(2026, 3, 1))


def test_parametros_e_volta_ao_padrao(entrar, operador, carregado, db):
    cliente = entrar(operador)
    cliente.post(
        "/configuracao/4561/parametros",
        data={"tributos": "16,5", "taxa_adm": "10", "pis_cofins": "3,65"},
    )
    assert db.session.get(Parametros, "4561").tributos == Decimal("0.1650")
    assert db.session.get(Parametros, "4561").pis_cofins == Decimal("0.0365")
    cliente.post("/configuracao/4561/parametros", data={"padrao": "on"})
    assert db.session.get(Parametros, "4561") is None


# --- planilha ---------------------------------------------------------------


def test_modelo_em_branco_tem_as_abas(entrar, operador, carregado):
    resposta = entrar(operador).get("/configuracao/modelo.xlsx")
    livro = load_workbook(io.BytesIO(resposta.data))
    assert {"Instruções", "Contrato", "Itens", "De-para", "BMs", "Pendências", "Pleitos"} <= set(
        livro.sheetnames
    )
    assert livro["Itens"].max_row == 1


def test_ida_e_volta_sem_perda(entrar, operador, carregado, db):
    cliente = entrar(operador)
    planilha = cliente.get("/configuracao/4561/planilha.xlsx").data
    previa = cliente.post(
        "/configuracao/importar",
        data={"arquivo": (io.BytesIO(planilha), "c.xlsx")},
        content_type="multipart/form-data",
    ).get_data(as_text=True)
    assert "Prévia da importação" in previa and "erro(s)" not in previa
    import re

    sha = re.search(r'name="sha" value="([0-9a-f]{64})"', previa).group(1)
    cliente.post("/configuracao/importar/confirmar", data={"sha": sha, "nome": "c.xlsx"})
    assert len(db.session.scalars(select(CfgItem).where(CfgItem.cr_norm == "4561")).all()) == 2
    assert len(db.session.scalars(select(CfgPleito).where(CfgPleito.cr_norm == "4561")).all()) == 2


def _planilha(linhas_por_aba: dict) -> bytes:
    from openpyxl import Workbook

    from app.configuracao.planilha import ABAS

    livro = Workbook()
    livro.remove(livro.active)
    for aba, linhas in linhas_por_aba.items():
        folha = livro.create_sheet(aba)
        folha.append(ABAS[aba])
        for linha in linhas:
            folha.append(linha)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _importar(cliente, conteudo: bytes) -> str:
    return cliente.post(
        "/configuracao/importar",
        data={"arquivo": (io.BytesIO(conteudo), "p.xlsx")},
        content_type="multipart/form-data",
    ).get_data(as_text=True)


def test_previa_conta_e_nao_grava(entrar, operador, carregado, db):
    corpo = _importar(
        entrar(operador),
        _planilha({"Itens": [["4602", "1.1", "Fiscal", "Mensal", 10, 1000, None, None]]}),
    )
    assert "Itens" in corpo and "Nada foi gravado ainda" in corpo
    assert db.session.scalars(select(CfgItem).where(CfgItem.cr_norm == "4602")).first() is None


def test_erros_por_linha_barram_a_confirmacao(entrar, operador, carregado):
    corpo = _importar(
        entrar(operador),
        _planilha(
            {
                "Itens": [["4561", "1.1", "", None, None, 10, None, None]],
                "Pleitos": [["9999", "03/2026", "x", "ADITIVO", "FEITO", 1, None]],
            }
        ),
    )
    assert "Itens, linha 2" in corpo and "CR 9999 não tem contrato" in corpo
    assert (
        "disabled>Confirmar importação" in corpo
        or "disabled" in corpo.split("Confirmar importação")[0][-80:]
    )


def test_cr_so_do_controle_na_planilha_e_recusado(entrar, operador, carregado):
    corpo = _importar(
        entrar(operador), _planilha({"Itens": [["4655", "1.1", "x", None, None, 10, None, None]]})
    )
    assert "não tem contrato na Receita" in corpo


def test_nf_que_nao_existe_vira_aviso(entrar, operador, carregado):
    corpo = _importar(
        entrar(operador), _planilha({"BMs": [["4561", "999", None, None, 3, "01/2026"]]})
    )
    assert "NF 999 ainda não chegou da Receita" in corpo


def test_planilha_nao_toca_em_contrato(entrar, operador, carregado, db):
    antes = db.session.scalars(select(RcContrato)).all()
    _importar(entrar(operador), _planilha({"Contrato": [["4561", "x", "31/12/2027", 20, 15]]}))
    assert db.session.scalars(select(RcContrato)).all() == antes


def test_depara_marca_natureza_que_gera_credito(entrar, operador, carregado, db):
    from app.models import CfgNaturezaCredito

    cliente = entrar(operador)
    dados = {
        "dp-0-natureza": "LOCACAO DE VEICULOS",
        "dp-0-item": "1.2",
        "dp-0-credito": "1",
        "dp-1-natureza": "HONORARIOS PJ - ENGENHARIA",
        "dp-1-item": "1.1",
    }
    assert cliente.post("/configuracao/4561/depara", data=dados).status_code == 302
    assert db.session.scalars(select(CfgNaturezaCredito.natureza_nome_norm)).all() == [
        "LOCACAO DE VEICULOS"
    ]
    corpo = cliente.get("/configuracao/?cr=4602&aba=depara").get_data(as_text=True)
    assert "Gera crédito PIS/COFINS" in corpo
    del dados["dp-0-credito"]
    cliente.post("/configuracao/4561/depara", data=dados)
    assert db.session.scalars(select(CfgNaturezaCredito)).all() == []


def test_parametros_globais_valem_para_quem_nao_tem_proprio(entrar, operador, carregado, db):
    from app.analise.dados import carregar_contratos

    cliente = entrar(operador)
    assert "Aplicar em" in cliente.get("/configuracao/globais").get_data(as_text=True)
    cliente.post(
        "/configuracao/4602/parametros",
        data={"tributos": "20", "taxa_adm": "10", "pis_cofins": "9,25"},
    )
    resposta = cliente.post(
        "/configuracao/globais",
        data={"alvo": "todos", "tributos": "12", "taxa_adm": "18", "pis_cofins": "3,65"},
    )
    assert resposta.status_code == 302
    assert db.session.get(Parametros, "*").pis_cofins == Decimal("0.0365")
    por_cr = {c.cr: c for c in carregar_contratos(db.session)}
    assert (por_cr["4561"].tributos, por_cr["4561"].taxa_adm) == (0.12, 0.18)
    assert por_cr["4561"].pis_cofins == 0.0365
    assert (por_cr["4602"].taxa_adm, por_cr["4602"].pis_cofins) == (0.10, 0.0925)


def test_global_substitui_os_proprios_quando_pedido(entrar, operador, carregado, db):
    from app.analise.dados import carregar_contratos

    cliente = entrar(operador)
    cliente.post(
        "/configuracao/4602/parametros",
        data={"tributos": "20", "taxa_adm": "10", "pis_cofins": "9,25"},
    )
    cliente.post(
        "/configuracao/globais", data={"alvo": "todos", "taxa_adm": "18", "sobrescrever": "1"}
    )
    por_cr = {c.cr: c for c in carregar_contratos(db.session)}
    assert por_cr["4602"].taxa_adm == 0.18 and por_cr["4602"].tributos == 0.20


def test_aplicar_por_cliente_so_nos_contratos_dele(entrar, operador, carregado, db):
    from app.analise.dados import carregar_contratos

    cliente = entrar(operador)
    cliente.post(
        "/configuracao/globais",
        data={"alvo": "cli", "sel": "Prefeitura Municipal", "pis_cofins": "3,65"},
    )
    por_cr = {c.cr: c for c in carregar_contratos(db.session)}
    assert por_cr["4602"].pis_cofins == 0.0365 and por_cr["4660"].pis_cofins == 0.0365
    assert por_cr["4561"].pis_cofins == 0.0925
    assert por_cr["4602"].taxa_adm == 0.15  # em branco: fica como estava
    assert db.session.get(Parametros, "*") is None


def test_aplicar_por_coordenador_e_por_centro_de_custo(entrar, operador, carregado, db):
    from app.analise.dados import carregar_contratos

    cliente = entrar(operador)
    cliente.post(
        "/configuracao/globais", data={"alvo": "coord", "sel": "Carlos Menezes", "taxa_adm": "12"}
    )
    cliente.post("/configuracao/globais", data={"alvo": "cr", "sel": "4660", "tributos": "11"})
    por_cr = {c.cr: c for c in carregar_contratos(db.session)}
    assert [por_cr[cr].taxa_adm for cr in ("4561", "4602", "4660")] == [0.12, 0.12, 0.15]
    assert [por_cr[cr].tributos for cr in ("4561", "4602", "4660")] == [0.20, 0.20, 0.11]


def test_aplicar_sem_selecao_nao_salva(entrar, operador, carregado, db):
    cliente = entrar(operador)
    cliente.post("/configuracao/globais", data={"alvo": "cli", "taxa_adm": "12"})
    cliente.post("/configuracao/globais", data={"alvo": "todos"})
    assert db.session.scalars(select(Parametros)).all() == []
